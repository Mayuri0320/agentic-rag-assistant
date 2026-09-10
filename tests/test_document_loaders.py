"""Tests for real document loaders."""

from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.ingestion.loaders.docx import DocxLoader
from app.ingestion.loaders.factory import DocumentLoaderFactory
from app.ingestion.loaders.pdf import PDFLoader
from app.ingestion.loaders.text import TextLoader


def test_text_loader(tmp_path: Path) -> None:
    """Text loader should read UTF-8 content."""
    file_path = tmp_path / "notes.txt"
    file_path.write_text(
        "Agentic RAG Assistant\nKnowledge retrieval",
        encoding="utf-8",
    )

    result = TextLoader().load(file_path)

    assert result == "Agentic RAG Assistant\nKnowledge retrieval"


def test_docx_loader(tmp_path: Path) -> None:
    """DOCX loader should extract paragraph text."""
    file_path = tmp_path / "document.docx"

    document = DocxDocument()
    document.add_paragraph("Agentic RAG Assistant")
    document.add_paragraph("Document ingestion pipeline")
    document.save(file_path)

    result = DocxLoader().load(file_path)

    assert "Agentic RAG Assistant" in result
    assert "Document ingestion pipeline" in result


def test_pdf_loader_empty_pdf(tmp_path: Path) -> None:
    """PDF loader should handle a PDF containing no extractable text."""
    file_path = tmp_path / "empty.pdf"

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    with file_path.open("wb") as file:
        writer.write(file)

    result = PDFLoader().load(file_path)

    assert result == ""


def test_default_factory_contains_real_loaders() -> None:
    """Default factory should register all supported loaders."""
    factory = DocumentLoaderFactory.default()

    assert isinstance(
        factory.get_loader(
            file_type="application/pdf",
            filename="report.pdf",
        ),
        PDFLoader,
    )

    assert isinstance(
        factory.get_loader(
            file_type="text/plain",
            filename="notes.txt",
        ),
        TextLoader,
    )

    assert isinstance(
        factory.get_loader(
            file_type=(
                "application/"
                "vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            filename="paper.docx",
        ),
        DocxLoader,
    )
